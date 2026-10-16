"""Hirelune ATS analysis.

A transparent, rule-based resume check against a target job. It is Hirelune's
own heuristic — not the score of any specific company's applicant tracking
system — and every point in the score is explained in the report.

Score (0-100) = keywords 40 + sections 20 + contact 10 + formatting 15 + impact 15
"""

from __future__ import annotations

import logging
import re
from collections import Counter
from typing import Optional

from ..agents.skill_agent import ALIASES, normalize_skill, parse_requirements

# Common skills/keywords recognised in free-text job descriptions.
VOCABULARY = [
    "python", "java", "javascript", "typescript", "c++", "c#", "go", "rust", "kotlin", "swift", "php", "ruby", "scala",
    "sql", "mysql", "postgresql", "mongodb", "redis", "sqlite", "oracle", "nosql", "elasticsearch",
    "react", "angular", "vue", "next.js", "node", "express", "django", "flask", "fastapi", "spring", "spring boot",
    ".net", "html", "css", "tailwind", "redux", "graphql", "rest apis", "microservices",
    "aws", "azure", "gcp", "docker", "kubernetes", "terraform", "jenkins", "ci/cd", "git", "github", "linux",
    "machine learning", "deep learning", "nlp", "computer vision", "pytorch", "tensorflow", "keras", "scikit-learn",
    "pandas", "numpy", "matplotlib", "tableau", "power bi", "excel", "statistics", "data analysis",
    "data visualization", "etl", "spark", "hadoop", "airflow", "kafka", "snowflake", "llm", "langchain",
    "generative ai", "artificial intelligence", "data structures", "algorithms", "oop", "system design",
    "unit testing", "selenium", "jira", "agile", "scrum", "figma", "communication", "leadership", "teamwork",
    "problem solving", "stakeholder management", "project management",
]

SECTIONS = {
    "summary": r"(summary|profile|objective|about me)",
    "experience": r"(experience|employment|work history|internships?)",
    "education": r"(education|academic|qualifications?)",
    "skills": r"(skills|technical skills|technologies|tech stack|competencies)",
    "projects": r"(projects?|portfolio)",
    "certifications": r"(certifications?|certificates|licenses|courses)",
}
REQUIRED_SECTIONS = ("experience", "education", "skills")
ACTION_VERBS = {
    "achieved", "analyzed", "analysed", "automated", "built", "created", "delivered", "designed", "developed",
    "drove", "improved", "implemented", "increased", "launched", "led", "managed", "optimized", "optimised",
    "reduced", "resolved", "shipped", "streamlined", "trained", "deployed", "migrated", "owned", "mentored",
}


logger = logging.getLogger(__name__)


def extract_text(path: str) -> str:
    """Prefer the platform's resume parser (handles two-column PDFs); fall back
    to a plain pdfplumber / python-docx read if it is unavailable or fails."""
    try:
        from ..agents import resume_parser
        text = resume_parser.extract_text(path)
        if text and text.strip():
            return text
    except Exception:  # noqa: BLE001 - fall back to the basic reader below
        logger.warning("resume_parser extraction failed for ATS; using basic reader", exc_info=True)
    return _basic_extract_text(path)


def section_details(text: str) -> dict:
    """Content of the Skills / Education / Experience / Projects / Certifications
    sections, using the same section splitter as resume screening."""
    try:
        from ..agents import resume_parser as rp
        sections, _, _ = rp.split_sections(text)
    except Exception:  # noqa: BLE001
        logger.warning("Section split failed for ATS", exc_info=True)
        return {}
    out: dict = {}
    def safe(fn, default):
        try:
            return fn()
        except Exception:  # noqa: BLE001
            return default
    if sections.get("skills"):
        out["skills"] = safe(lambda: rp.extract_list_section(sections["skills"], 30, split_commas=True), [])
    if sections.get("education"):
        out["education"] = safe(lambda: rp.extract_list_section(sections["education"], 6), [])
    if sections.get("experience"):
        exp = sections["experience"]
        out["experience"] = safe(lambda: rp.extract_list_section(exp, 8), [])
        out["experience_years"] = safe(lambda: round(rp.merge_and_sum_years(rp.find_date_ranges(exp)), 1), None)
    if sections.get("projects"):
        out["projects"] = safe(lambda: [p["title"] for p in rp.extract_projects(sections["projects"])], [])
    if sections.get("certifications"):
        out["certifications"] = safe(lambda: [c["name"] for c in rp.extract_certifications(sections["certifications"])], [])
    return out


def _basic_extract_text(path: str) -> str:
    lower = path.lower()
    if lower.endswith(".pdf"):
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return "\n".join((page.extract_text() or "") for page in pdf.pages)
    if lower.endswith(".docx"):
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    raise ValueError("Only PDF and DOCX resumes can be analysed")


def _has_phrase(text_norm: str, phrase: str) -> bool:
    phrase = _normalized_text(phrase).strip()
    if not phrase:
        return False
    return re.search(rf"(?<![\w+#]){re.escape(phrase)}(?![\w+#])", text_norm) is not None


def _normalized_text(text: str) -> str:
    """Normalize the whole resume the same way skills are normalized, token-wise."""
    lowered = text.lower()
    lowered = re.sub(r"(?<=\w)\.js\b", "", lowered)
    lowered = re.sub(r"[^\w\s+#]|_", " ", lowered)
    lowered = re.sub(r"\s+", " ", lowered)
    # Apply multi-word/alias canonical forms by checking vocabulary aliases later.
    return f" {lowered.strip()} "


def job_keywords(job_description: str = "", requirements: str = "") -> list[str]:
    keywords: list[str] = []
    seen = set()
    for req in parse_requirements(requirements or ""):
        key = normalize_skill(req)
        if key and key not in seen:
            seen.add(key)
            keywords.append(req)
    jd_norm = _normalized_text(job_description or "")
    for term in VOCABULARY:
        key = normalize_skill(term)
        if key not in seen and (_has_phrase(jd_norm, term) or _has_phrase(jd_norm, key)):
            seen.add(key)
            keywords.append(term)
    return keywords


def _match_terms(norm: str, terms: list[str]) -> tuple[list[str], list[str]]:
    matched, missing = [], []
    for kw in terms:
        key = normalize_skill(kw)
        aliases = {kw.lower(), key} | {a for a, canon in ALIASES.items() if canon == key}
        (matched if any(_has_phrase(norm, a) for a in aliases) else missing).append(kw)
    return matched, missing


_DEGREE_LEVELS = [  # (level, pattern) - higher level satisfies lower requirements
    (1, r"diploma"),
    (2, r"b\.?\s?tech|b\.?\s?e\b|bachelor|b\.?\s?sc|bca|b\.?\s?com|b\.?\s?a\b|undergraduate|graduate"),
    (3, r"m\.?\s?tech|master|m\.?\s?sc|mca|mba|m\.?\s?e\b|post\s?graduate"),
    (4, r"ph\.?\s?d|doctorate"),
]


def _degree_level(text: str) -> int:
    return max((lvl for lvl, pat in _DEGREE_LEVELS if re.search(rf"\b({pat})", text or "", re.I)), default=0)


def analyze(resume_text: str, job_description: str = "", requirements: str = "",
            min_experience: Optional[float] = None, preferred_skills: str = "",
            education_requirement: str = "", sections_detail: Optional[dict] = None) -> dict:
    text = resume_text or ""
    words = re.findall(r"[A-Za-z][A-Za-z+#.\-]*", text)
    word_count = len(words)
    norm = _normalized_text(text)

    # ---- Keywords (40) ----
    keywords = job_keywords(job_description, requirements)
    matched, missing = [], []
    for kw in keywords:
        key = normalize_skill(kw)
        aliases = {kw.lower(), key}
        aliases |= {a for a, canon in ALIASES.items() if canon == key}
        (matched if any(_has_phrase(norm, a) for a in aliases) else missing).append(kw)
    keyword_score = round(40 * len(matched) / len(keywords), 1) if keywords else 0.0

    # ---- Sections (20) ----
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    found_sections = []
    for name, pattern in SECTIONS.items():
        if any(len(ln) <= 40 and re.fullmatch(rf"[\W\d_]*{pattern}[\W_]*", ln, flags=re.I) for ln in lines):
            found_sections.append(name)
    missing_sections = [s for s in SECTIONS if s not in found_sections]
    required_found = sum(1 for s in REQUIRED_SECTIONS if s in found_sections)
    optional_found = sum(1 for s in found_sections if s not in REQUIRED_SECTIONS)
    section_score = round(min(20, required_found * 5 + optional_found * 5 / 3), 1)

    # ---- Contact (10) ----
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    phone = re.search(r"(\+?\d[\d\s().-]{8,}\d)", text)
    linkedin = re.search(r"linkedin\.com/\S+", text, re.I)
    github = re.search(r"(github\.com|gitlab\.com)/\S+", text, re.I)
    contact_score = (4 if email else 0) + (3 if phone else 0) + (3 if (linkedin or github) else 0)

    # ---- Formatting (15) ----
    warnings = []
    formatting = 15.0
    if word_count < 150:
        warnings.append("Very little text could be read. If the resume is a scanned image or uses graphics for text, "
                        "ATS software may not be able to read it — use a text-based PDF or DOCX.")
        formatting -= 8
    if word_count > 1200:
        warnings.append(f"The resume is long ({word_count} words). Aim for 1–2 pages (roughly 400–900 words).")
        formatting -= 3
    odd = len(re.findall(r"[^\x00-\x7F₹€£•–—’“”‘]", text))
    if word_count and odd / max(1, len(text)) > 0.02:
        warnings.append("Unusual symbols or icons were detected; some ATS parsers drop or garble them.")
        formatting -= 2
    if re.search(r"\|.*\|.*\|", text):
        warnings.append("Table-like layouts (columns separated by |) were detected; simple single-column layouts parse best.")
        formatting -= 2
    dates = re.findall(r"((jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s+\d{4}|\b(19|20)\d{2}\b)", text, re.I)
    if "experience" in found_sections and len(dates) < 2:
        warnings.append("Few dates found. Add start and end dates (e.g. Jan 2023 – Present) to each role.")
        formatting -= 2
    if not email:
        warnings.append("No email address found in the text.")
    formatting = max(0.0, formatting)

    # ---- Impact (15) ----
    lower_words = [w.lower().strip(".-") for w in words]
    verbs = sorted({w for w in lower_words if w in ACTION_VERBS})
    numbers = []
    for ln in lines:
        if "@" in ln or re.search(r"\+?\d[\d\s().-]{8,}\d", ln):
            continue  # contact line
        for m in re.finditer(r"(?<![\w.])(\d+(?:\.\d+)?)\s?(%|percent|x\b|k\b|\+)?", ln):
            value = m.group(1)
            if m.group(2) or not re.fullmatch(r"(19|20)\d{2}", value):
                numbers.append(value)
    impact_score = min(8, len(verbs) * 1.5) + min(7, len(numbers) * 2)

    score = round(keyword_score + section_score + contact_score + formatting + impact_score, 1)

    # ---- Relevance notes ----
    years = [float(y) for y in re.findall(r"(\d+(?:\.\d+)?)\+?\s*(?:years?|yrs?)", text, re.I)]
    experience_note = None
    if min_experience:
        best = max(years) if years else None
        if best is None:
            experience_note = f"The job asks for {min_experience:g}+ years; state your total experience explicitly."
        elif best < min_experience:
            experience_note = f"The resume mentions {best:g} years; the job asks for {min_experience:g}+."
        else:
            experience_note = f"Mentions {best:g} years, meeting the {min_experience:g}-year requirement."
    has_degree = re.search(r"\b(b\.?\s?tech|b\.?\s?e\b|bachelor|master|m\.?\s?tech|mba|b\.?\s?sc|m\.?\s?sc|bca|mca|ph\.?d|diploma)", text, re.I)

    # Preferred (nice-to-have) skills: reported, not part of the 0-100 score.
    pref_terms = [t for t in parse_requirements(preferred_skills or "")
                  if normalize_skill(t) not in {normalize_skill(k) for k in keywords}]
    pref_matched, pref_missing = _match_terms(norm, pref_terms)

    detail = sections_detail or {}
    parsed_years = detail.get("experience_years")
    stated_years = max(years) if years else None
    best_years = max([y for y in (parsed_years, stated_years) if y is not None], default=None)
    experience_match = None
    if min_experience is not None:
        experience_match = {"required": min_experience, "found": best_years,
                            "meets": (best_years or 0) >= min_experience}
    education_match = None
    if (education_requirement or "").strip():
        need, have = _degree_level(education_requirement), _degree_level(text)
        education_match = {"required": education_requirement.strip(), "found": bool(has_degree),
                           "meets": (have >= need) if need else bool(has_degree)}

    suggestions = []
    if missing:
        suggestions.append(f"Add evidence for missing keywords you genuinely have: {', '.join(missing[:8])}.")
    for s in REQUIRED_SECTIONS:
        if s not in found_sections:
            suggestions.append(f"Add a clearly titled '{s.title()}' section.")
    if len(numbers) < 3:
        suggestions.append("Quantify achievements (%, time saved, users, revenue) in at least three bullet points.")
    if len(verbs) < 5:
        suggestions.append("Start bullet points with strong action verbs such as built, improved, led, automated.")
    if not (linkedin or github):
        suggestions.append("Add a LinkedIn or GitHub profile link.")
    if not keywords:
        suggestions.append("Choose a job or paste a job description to get keyword matching.")
    if pref_missing:
        suggestions.append(f"Nice-to-have skills for this job you could highlight if you have them: {', '.join(pref_missing[:6])}.")
    if experience_match and not experience_match["meets"]:
        suggestions.append(experience_note or f"The job asks for {min_experience:g}+ years of experience.")
    if education_match and not education_match["meets"]:
        suggestions.append(f"The job lists '{education_match['required']}'. Make your degree and field easy to find in the Education section.")
    if "projects" not in found_sections:
        suggestions.append("Add a Projects section with 2–3 relevant projects and the tools used.")

    strengths = []
    if keywords and len(matched) / len(keywords) >= 0.7:
        strengths.append(f"Strong keyword coverage ({len(matched)} of {len(keywords)}).")
    if required_found == len(REQUIRED_SECTIONS):
        strengths.append("All core sections (experience, education, skills) are present.")
    if len(numbers) >= 3:
        strengths.append("Achievements are quantified.")
    if len(verbs) >= 5:
        strengths.append("Uses strong action verbs.")
    if email and phone:
        strengths.append("Contact details are easy to find.")

    return {
        "label": "Hirelune ATS analysis",
        "disclaimer": "Hirelune's own rule-based estimate of how well an applicant tracking system can read and match "
                      "this resume. It is not the score of any specific employer's ATS.",
        "score": min(100.0, score),
        "breakdown": {
            "keywords": {"score": keyword_score, "max": 40},
            "sections": {"score": section_score, "max": 20},
            "contact": {"score": contact_score, "max": 10},
            "formatting": {"score": round(formatting, 1), "max": 15},
            "impact": {"score": round(impact_score, 1), "max": 15},
        },
        "keywords": {"total": len(keywords), "matched": matched, "missing": missing},
        "sections": {"found": found_sections, "missing": missing_sections},
        "contact": {"email": bool(email), "phone": bool(phone), "linkedin": bool(linkedin), "github": bool(github)},
        "formatting_issues": warnings,
        "action_verbs": verbs,
        "quantified_achievements": len(numbers),
        "word_count": word_count,
        "experience_relevance": experience_note,
        "experience_match": experience_match,
        "education_match": education_match,
        "preferred_skills": {"total": len(pref_terms), "matched": pref_matched, "missing": pref_missing},
        "section_details": detail,
        "education_found": bool(has_degree),
        "strengths": strengths,
        "suggestions": suggestions,
        "top_terms": [w for w, _ in Counter(w for w in lower_words if len(w) > 3).most_common(10)],
    }
